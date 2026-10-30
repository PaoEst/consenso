# -*- coding: utf-8 -*-
"""Lectura de Consenso.xlsx y utilidades compartidas por los scripts de consenso."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
import plotly.offline

from graficar_aportes import BASE, log, norm

ARCHIVO_CONSENSO = "Consenso.xlsx"
ARCHIVO_PRONOSTICO = "Pronostico_Drivers.xlsx"
HOJA_FECHAS = "Fechas"
HOJA_EVALUADORES = "Evaluadores_Tecnicos"
HOJA_REGISTROS = "Registros"
N_SEMANAS = 5

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def fecha_corta(f: pd.Timestamp) -> str:
    return f"{f.day:02d} {MESES[f.month - 1][:3]}"


def etiqueta_semana_corta(s: dict) -> str:
    """'01-07 nov' o '28 oct - 03 nov' para los rótulos de semana."""
    if s["ini"].month == s["fin"].month:
        return f"{s['ini'].day:02d}-{fecha_corta(s['fin'])}"
    return f"{fecha_corta(s['ini'])} - {fecha_corta(s['fin'])}"


def _fecha(v):
    f = pd.to_datetime(v, errors="coerce", dayfirst=isinstance(v, str))
    return None if pd.isna(f) else f.normalize()


_TERMINO = re.compile(r"([+-]?)([A-Za-z]{1,2}[0-9]+|[0-9]+(?:\.[0-9]+)?)")


def _valor_celda(ws_f, ws_v, coord: str, profundidad: int = 0):
    """Valor de una celda; si no tiene valor guardado (Excel no recalculado, p. ej. tras guardar
    con openpyxl) evalúa fórmulas simples del tipo =+C2, =+E2+6, =F2-1."""
    v = ws_v[coord].value
    if v is not None:
        return v
    f = ws_f[coord].value
    if not (isinstance(f, str) and f.startswith("=")) or profundidad > 20:
        return f
    expr = f[1:].replace(" ", "")
    total, pos, fecha = 0.0, 0, False
    while pos < len(expr):
        m = _TERMINO.match(expr, pos)
        if not m:
            raise ValueError(f"Fórmula no soportada en {coord}: {f}")
        signo = -1 if m.group(1) == "-" else 1
        tok = m.group(2)
        if tok[0].isalpha():
            ref = _valor_celda(ws_f, ws_v, tok.upper(), profundidad + 1)
            if isinstance(ref, (datetime, date)):
                fecha = True
                ref = (ref.date() if isinstance(ref, datetime) else ref).toordinal()
            if ref is None:
                return None
            total += signo * float(ref)
        else:
            total += signo * float(tok)
        pos = m.end()
    return datetime.fromordinal(int(round(total))) if fecha else total


def leer_fechas(ruta: Path) -> dict:
    """Hoja Fechas (fila 2): consenso, horizonte y semanas S1..S5 (inicio/fin)."""
    wb_f = openpyxl.load_workbook(ruta)
    wb_v = openpyxl.load_workbook(ruta, data_only=True)
    if HOJA_FECHAS not in wb_f.sheetnames:
        raise ValueError(f"{ruta.name} no tiene la hoja '{HOJA_FECHAS}'")
    ws_f, ws_v = wb_f[HOJA_FECHAS], wb_v[HOJA_FECHAS]
    v = [_valor_celda(ws_f, ws_v, f"{get_column_letter(c)}2") for c in range(1, 5 + 2 * N_SEMANAS)]
    out = {"seguimiento_ini": _fecha(v[0]), "consenso": _fecha(v[1]),
           "hor_ini": _fecha(v[2]), "hor_fin": _fecha(v[3]), "semanas": []}
    for faltante in ("consenso", "hor_ini", "hor_fin"):
        if out[faltante] is None:
            raise ValueError(f"Hoja '{HOJA_FECHAS}': falta o es ilegible la fecha '{faltante}'")
    if out["hor_ini"] > out["hor_fin"]:
        raise ValueError("Hoja Fechas: el inicio del horizonte es posterior a su fin")
    for k in range(N_SEMANAS):
        ini, fin = _fecha(v[4 + 2 * k]), _fecha(v[5 + 2 * k])
        if ini is None or fin is None:
            log.info("Fechas: S%d sin fechas completas; no se usa", k + 1)
            continue
        if ini > fin:
            raise ValueError(f"Hoja Fechas: S{k + 1} tiene inicio posterior al fin")
        if ini < out["hor_ini"] or fin > out["hor_fin"]:
            log.warning("Fechas: S%d (%s -> %s) cae fuera del horizonte (%s -> %s)", k + 1,
                        ini.date(), fin.date(), out["hor_ini"].date(), out["hor_fin"].date())
        out["semanas"].append({"n": k + 1, "ini": ini, "fin": fin, "dias": (fin - ini).days + 1})
    if not out["semanas"]:
        raise ValueError("Hoja Fechas: no hay ninguna semana definida")
    log.info("Fechas: consenso=%s horizonte=%s -> %s semanas=%s", out["consenso"].date(),
             out["hor_ini"].date(), out["hor_fin"].date(),
             [(s["n"], str(s["ini"].date()), str(s["fin"].date()), s["dias"]) for s in out["semanas"]])
    return out


def leer_evaluadores(ruta: Path) -> dict[str, str]:
    """{INICIALES_MAYUSCULA: nombre} de la hoja Evaluadores_Tecnicos."""
    df = pd.read_excel(ruta, sheet_name=HOJA_EVALUADORES)
    df = df.iloc[:, :2].dropna(how="all")
    df.columns = ["nombre", "iniciales"]
    out = {}
    for _, r in df.dropna(subset=["iniciales"]).iterrows():
        ini = str(r["iniciales"]).strip().upper()
        if ini in out:
            log.warning("Evaluadores: iniciales repetidas '%s'; se conserva la primera", ini)
            continue
        out[ini] = str(r["nombre"]).strip()
    log.info("Evaluadores técnicos (%d): %s", len(out), out)
    return out


def carpeta_corrida(base: Path, fecha: pd.Timestamp) -> Path:
    """<base>/<año>/<mes>/<AAAA-MM-DD>"""
    return base / f"{fecha.year}" / f"{fecha.month:02d}" / f"{fecha:%Y-%m-%d}"


def plotly_js() -> str:
    """Código de plotly.js para incrustarlo en el HTML (funciona sin internet)."""
    return plotly.offline.get_plotlyjs()

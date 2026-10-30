#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Registra los votos de los evaluadores en Consenso.xlsx > Registros y genera, en la misma
corrida, el gráfico de consolidación por cuenca y por ronda.

Flujo
  1. Cada evaluador abre su Consenso_<cuenca>.html, vota y pulsa GUARDAR (se descarga un
     archivo voto_*.json). Los archivos se dejan en  <año>/<mes>/<fecha>/votos/
  2. Se corre este script. Para cada voto nuevo agrega una fila a Registros (sin duplicar) y
     luego crea  <año>/<mes>/<fecha>/consolidacion/Consolidacion_<cuenca>_Ronda<N>.html

Rondas: el "No Evaluación" es la ronda. En la ronda N cada evaluador cuenta con su voto más
reciente de ronda <= N (si no volvió a votar, se mantiene el de la ronda anterior).

Uso:  python consolidar_votos.py [--ronda N] [--fecha-corrida AAAA-MM-DD]
"""
from __future__ import annotations

import argparse
import getpass
import html
import json
import os
import platform
import shutil
import statistics
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import plotly
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from consenso_comun import (ARCHIVO_CONSENSO, BASE, HOJA_REGISTROS, N_SEMANAS, carpeta_corrida,
                            etiqueta_semana_corta, leer_evaluadores, leer_fechas, plotly_js)
from graficar_aportes import (ARCHIVO_APORTES, C_ABAJO, C_ARRIBA, C_LINEA, configurar_logging,
                              definir_ventana, leer_aportes, log, nombre_archivo, norm, relleno_media,
                              sha256)

COLUMNAS_BASE = ["Fecha_hora_registro", "No Evaluación", "Cuenca", "Evaluador", "Aportes S1",
                 "Aportes S2", "Aportes S3", "Aportes S4", "Aportes S5", "Real a la fecha",
                 "Aportes_resto del mes", "Aportes cierre de mes"]
COLUMNAS_EXTRA = ["Mes_pronostico", "Id_voto"]
PALETA = ["#1f77b4", "#ff7f0e", "#2ca02c", "#9467bd", "#8c564b", "#e377c2", "#17becf", "#bcbd22"]


# --------------------------------------------------------------------------
# Lectura y validación de votos
# --------------------------------------------------------------------------
def _num(v):
    if v is None or v == "":
        return None
    f = float(v)
    if f != f:  # NaN
        raise ValueError("valor no numérico")
    return f


def leer_votos(carpeta_votos: Path, evaluadores: dict, cuencas: set[str], mes: str):
    """Devuelve (votos_validos, rechazados). Cada voto es un dict normalizado."""
    validos, rechazados, ids = [], [], set()
    archivos = sorted(carpeta_votos.glob("*.json"))
    log.info("Votos: %d archivos .json en %s", len(archivos), carpeta_votos)
    for f in archivos:
        try:
            d = json.loads(f.read_text(encoding="utf-8-sig"))
            vid = str(d["id"]).strip()
            ev = str(d["evaluador"]).strip().upper()
            if ev not in evaluadores:
                raise ValueError(f"evaluador '{ev}' no está en Evaluadores_Tecnicos")
            ronda = int(d["no_evaluacion"])
            if ronda < 1 or ronda != float(d["no_evaluacion"]):
                raise ValueError("No Evaluación debe ser un entero >= 1")
            cuenca = str(d["cuenca"]).strip()
            if cuenca not in cuencas:
                raise ValueError(f"cuenca '{cuenca}' no existe en Seguimiento_Aportes")
            if str(d["mes_pronostico"]) != mes:
                raise ValueError(f"el voto es del mes {d['mes_pronostico']}, la corrida es de {mes}")
            sem = {f"S{k}": _num(d["semanas"].get(f"S{k}")) for k in range(1, N_SEMANAS + 1)}
            if any(v is not None and v < 0 for v in sem.values()):
                raise ValueError("hay aportes negativos")
            if all(v is None for v in sem.values()):
                raise ValueError("el voto no tiene ninguna semana")
            resto, cierre = _num(d["resto_del_mes"]), _num(d["cierre_de_mes"])
            if resto is None or cierre is None:
                raise ValueError("faltan resto del mes o cierre de mes")
            if vid in ids:
                raise ValueError(f"id de voto repetido en otro archivo ({vid})")
            ids.add(vid)
            validos.append({
                "id": vid, "fecha_hora": datetime.fromisoformat(str(d["fecha_hora"])),
                "evaluador": ev, "ronda": ronda, "cuenca": cuenca, "semanas": sem,
                "real": _num(d.get("real_a_la_fecha")), "resto": resto, "cierre": cierre,
                "mes": mes, "archivo": f.name})
            log.debug("Voto válido: %s | %s | %s | ronda %d | %s", f.name, ev, cuenca, ronda, sem)
        except Exception as e:
            rechazados.append((f.name, str(e)))
            log.warning("Voto RECHAZADO %s: %s", f.name, e)
    log.info("Votos válidos=%d rechazados=%d", len(validos), len(rechazados))
    return validos, rechazados


# --------------------------------------------------------------------------
# Registros (Excel)
# --------------------------------------------------------------------------
def cargar_registros(ruta: Path):
    wb = openpyxl.load_workbook(ruta)
    if HOJA_REGISTROS not in wb.sheetnames:
        raise ValueError(f"{ruta.name} no tiene la hoja '{HOJA_REGISTROS}'")
    ws = wb[HOJA_REGISTROS]
    cols = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v is not None:
            cols[norm(v)] = c
    for extra in COLUMNAS_EXTRA:
        if norm(extra) not in cols:
            c = max(cols.values()) + 1
            ws.cell(1, c, extra)
            cols[norm(extra)] = c
            log.info("Registros: se agregó la columna '%s' en la posición %d", extra, c)
    faltan = [c for c in COLUMNAS_BASE if norm(c) not in cols]
    if faltan:
        raise ValueError(f"Registros: faltan las columnas {faltan}")
    return wb, ws, cols


def ultima_fila(ws, cols) -> int:
    """Última fila con algún dato en las columnas de Registros."""
    for r in range(ws.max_row, 1, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in cols.values()):
            return r
    return 1


def escribir_votos(ruta: Path, votos: list[dict], dir_logs: Path):
    """Agrega los votos nuevos a Registros. Devuelve (wb, ws, cols, nuevos)."""
    wb, ws, cols = cargar_registros(ruta)
    ya = {str(ws.cell(r, cols[norm("Id_voto")]).value) for r in range(2, ws.max_row + 1)
          if ws.cell(r, cols[norm("Id_voto")]).value}
    nuevos = [v for v in votos if v["id"] not in ya]
    log.info("Registros: %d votos ya estaban registrados, %d son nuevos", len(votos) - len(nuevos), len(nuevos))
    if not nuevos:
        return wb, ws, cols, []
    # copia de seguridad antes de escribir
    dir_bk = dir_logs / "backups"
    dir_bk.mkdir(parents=True, exist_ok=True)
    bk = dir_bk / f"{ruta.stem}_{datetime.now():%Y%m%d_%H%M%S}{ruta.suffix}"
    shutil.copy2(ruta, bk)
    log.info("Copia de seguridad: %s", bk.resolve())

    fila = ultima_fila(ws, cols) + 1
    for v in sorted(nuevos, key=lambda x: (x["fecha_hora"], x["id"])):
        vals = {"Fecha_hora_registro": v["fecha_hora"], "No Evaluación": v["ronda"],
                "Cuenca": v["cuenca"], "Evaluador": v["evaluador"],
                "Real a la fecha": v["real"], "Aportes_resto del mes": v["resto"],
                "Aportes cierre de mes": v["cierre"], "Mes_pronostico": v["mes"],
                "Id_voto": v["id"]}
        for k in range(1, N_SEMANAS + 1):
            vals[f"Aportes S{k}"] = v["semanas"][f"S{k}"]
        for nombre, val in vals.items():
            ws.cell(fila, cols[norm(nombre)], val)
        ws.cell(fila, cols[norm("Fecha_hora_registro")]).number_format = "yyyy-mm-dd hh:mm:ss"
        log.info("Registros fila %d <- %s | %s | ronda %d | resto=%s cierre=%s | id=%s", fila,
                 v["evaluador"], v["cuenca"], v["ronda"], v["resto"], v["cierre"], v["id"])
        fila += 1
    tmp = ruta.with_name(ruta.stem + ".tmp" + ruta.suffix)
    wb.save(tmp)
    try:
        os.replace(tmp, ruta)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise PermissionError(f"No se pudo escribir {ruta.name}: ciérralo en Excel y vuelve a correr "
                              f"(los votos NO se registraron; la copia {bk.name} no cambió nada).")
    return wb, ws, cols, nuevos


def votos_del_mes(ws, cols, mes: str) -> list[dict]:
    """Todos los votos del mes registrados en Registros (todas las corridas)."""
    g = lambda r, n: ws.cell(r, cols[norm(n)]).value
    out, omitidos = [], 0
    for r in range(2, ws.max_row + 1):
        cuenca, ev, ronda, m = g(r, "Cuenca"), g(r, "Evaluador"), g(r, "No Evaluación"), g(r, "Mes_pronostico")
        if not cuenca or not ev or ronda in (None, ""):
            omitidos += 1 if any(ws.cell(r, c).value not in (None, "") for c in cols.values()) else 0
            continue
        if str(m) != mes:
            continue
        fh = g(r, "Fecha_hora_registro")
        out.append({"cuenca": str(cuenca).strip(), "evaluador": str(ev).strip().upper(),
                    "ronda": int(ronda), "fecha_hora": fh if isinstance(fh, datetime) else datetime.min,
                    "semanas": {f"S{k}": g(r, f"Aportes S{k}") for k in range(1, N_SEMANAS + 1)},
                    "resto": g(r, "Aportes_resto del mes"), "cierre": g(r, "Aportes cierre de mes"),
                    "real": g(r, "Real a la fecha")})
    if omitidos:
        log.warning("Registros: %d filas incompletas o sin cuenca/mes se ignoran en la consolidación", omitidos)
    return out


# --------------------------------------------------------------------------
# Consolidación
# --------------------------------------------------------------------------
def votos_vigentes(votos_cuenca: list[dict], ronda: int) -> list[dict]:
    """Por evaluador, su voto más reciente con ronda <= ronda (arrastra rondas anteriores)."""
    por_ev = {}
    for v in sorted(votos_cuenca, key=lambda x: (x["ronda"], x["fecha_hora"])):
        if v["ronda"] <= ronda:
            if v["evaluador"] in por_ev and por_ev[v["evaluador"]]["ronda"] == v["ronda"]:
                log.warning("Evaluador %s votó más de una vez la ronda %d (%s); se usa el último",
                            v["evaluador"], v["ronda"], v["cuenca"])
            por_ev[v["evaluador"]] = v
    return sorted(por_ev.values(), key=lambda x: x["evaluador"])


def estadisticos(valores: list[float]) -> dict:
    v = [x for x in valores if x is not None]
    if not v:
        return {"media": None, "mediana": None, "min": None, "max": None}
    return {"media": sum(v) / len(v), "mediana": statistics.median(v), "min": min(v), "max": max(v)}


def r1(x):
    return None if x is None else round(float(x), 1)


def construir_consolidacion(cuenca, ronda, vigentes, historial, serie, fechas, evaluadores):
    hor_ini, hor_fin = fechas["hor_ini"], fechas["hor_fin"]
    x_fin = hor_fin + pd.Timedelta(days=1)
    semanas = fechas["semanas"]
    ini = serie["Fecha"].min() if len(serie) else hor_ini - pd.Timedelta(days=30)

    fig = make_subplots(rows=2, cols=1, vertical_spacing=0.12, row_heights=[0.66, 0.34],
                        subplot_titles=("Aportes: seguimiento y pronóstico por evaluador",
                                        "Esperado resto del mes y cierre de mes"))
    # Seguimiento
    if len(serie):
        x = list(serie["Fecha"] + pd.Timedelta(hours=12))
        y = list(serie[cuenca])
        xs, arr, abj = relleno_media(x, y)
        base = [100.0] * len(xs)
        for techo, color in ((arr, C_ARRIBA), (abj, C_ABAJO)):
            fig.add_trace(go.Scatter(x=xs, y=base, mode="lines", line=dict(width=0),
                                     hoverinfo="skip", showlegend=False), row=1, col=1)
            fig.add_trace(go.Scatter(x=xs, y=techo, mode="lines", line=dict(width=0), fill="tonexty",
                                     fillcolor=color, hoverinfo="skip", showlegend=False), row=1, col=1)
        fig.add_trace(go.Scatter(
            x=x, y=y, mode="lines", line=dict(color=C_LINEA, width=2.2), showlegend=False,
            hovertemplate="%{x|%d %b %Y}<br><b>%{y:.1f} %</b> de la media<extra></extra>"),
            row=1, col=1)
        fig.add_hline(y=100, line=dict(color="#333", width=1, dash="dash"), row=1, col=1)

    def escalon(valores):
        xs, ys = [], []
        for s in semanas:
            val = valores.get(f"S{s['n']}")
            if val is None:
                if xs and xs[-1] is not None:
                    xs.append(None)
                    ys.append(None)
                continue
            xs += [s["ini"], s["fin"] + pd.Timedelta(days=1)]
            ys += [val, val]
        return xs, ys

    # Una línea por evaluador
    colores = {}
    for i, v in enumerate(vigentes):
        colores[v["evaluador"]] = PALETA[i % len(PALETA)]
        xs, ys = escalon(v["semanas"])
        nombre = evaluadores.get(v["evaluador"], v["evaluador"])
        fig.add_trace(go.Scatter(
            x=xs, y=ys, mode="lines", name=f"{v['evaluador']} (R{v['ronda']})",
            line=dict(color=colores[v["evaluador"]], width=2),
            hovertemplate=f"<b>{html.escape(nombre)}</b> · ronda {v['ronda']}<br>%{{y:.1f}} %<extra></extra>"),
            row=1, col=1)

    # Estadísticos por semana
    est_sem = {f"S{s['n']}": estadisticos([v["semanas"].get(f"S{s['n']}") for v in vigentes])
               for s in semanas}
    grupos = {"media": [], "mediana": [], "rango": []}
    for clave, nombre, dash in (("media", "Media", "solid"), ("mediana", "Mediana", "dash")):
        xs, ys = escalon({k: e[clave] for k, e in est_sem.items()})
        grupos[clave].append(len(fig.data))
        fig.add_trace(go.Scatter(x=xs, y=ys, mode="lines", name=f"{nombre} del consenso",
                                 line=dict(color="#111", width=4, dash=dash),
                                 visible=(clave == "media"),
                                 hovertemplate=f"{nombre}: %{{y:.1f}} %<extra></extra>"), row=1, col=1)
    xs_min, ys_min = escalon({k: e["min"] for k, e in est_sem.items()})
    xs_max, ys_max = escalon({k: e["max"] for k, e in est_sem.items()})
    grupos["rango"].append(len(fig.data))
    fig.add_trace(go.Scatter(x=xs_min, y=ys_min, mode="lines", line=dict(color="#111", width=1),
                             name="Mínimo", visible=False, hovertemplate="Mínimo: %{y:.1f} %<extra></extra>"),
                  row=1, col=1)
    grupos["rango"].append(len(fig.data))
    fig.add_trace(go.Scatter(x=xs_max, y=ys_max, mode="lines", line=dict(color="#111", width=1),
                             fill="tonexty", fillcolor="rgba(17,17,17,0.15)", name="Máximo",
                             visible=False, hovertemplate="Máximo: %{y:.1f} %<extra></extra>"),
                  row=1, col=1)

    # Resto del mes y cierre: un punto por evaluador y marca del estadístico
    cat = {"Esperado resto del mes": 1, "Cierre de mes": 0}
    for v in vigentes:
        fig.add_trace(go.Scatter(
            x=[v["resto"], v["cierre"]], y=[1, 0], mode="markers", showlegend=False,
            marker=dict(color=colores[v["evaluador"]], size=13, line=dict(color="white", width=1.5)),
            hovertemplate=f"<b>{v['evaluador']}</b> (R{v['ronda']})<br>%{{x:.1f}} %<extra></extra>"),
            row=2, col=1)
    e_resto = estadisticos([v["resto"] for v in vigentes])
    e_cierre = estadisticos([v["cierre"] for v in vigentes])
    for clave, simbolo, nombre in (("media", "diamond", "Media"), ("mediana", "square", "Mediana")):
        grupos[clave].append(len(fig.data))
        fig.add_trace(go.Scatter(
            x=[e_resto[clave], e_cierre[clave]], y=[1, 0], mode="markers", showlegend=False,
            visible=(clave == "media"),
            marker=dict(color="#111", size=18, symbol=simbolo, line=dict(color="white", width=1.5)),
            hovertemplate=f"{nombre}: %{{x:.1f}} %<extra></extra>"), row=2, col=1)
    grupos["rango"].append(len(fig.data))
    fig.add_trace(go.Scatter(
        x=[e_resto["min"], e_resto["max"], None, e_cierre["min"], e_cierre["max"]],
        y=[1, 1, None, 0, 0], mode="lines", showlegend=False, visible=False,
        line=dict(color="#111", width=8), opacity=0.35, hoverinfo="skip"), row=2, col=1)

    formas = [dict(type="line", xref="x", yref="paper", x0=hor_ini, x1=hor_ini, y0=0.40, y1=1,
                   line=dict(color="#d62728", width=2, dash="dash"))]
    for s in semanas:
        if s["ini"] > hor_ini:
            formas.append(dict(type="line", xref="x", yref="paper", x0=s["ini"], x1=s["ini"],
                               y0=0.40, y1=1, line=dict(color="#2e9e4f", width=1, dash="dot")))
    fig.update_layout(
        height=760, template="plotly_white", shapes=formas, hovermode="closest",
        margin=dict(l=120, r=30, t=60, b=90), plot_bgcolor="#fbfbf9",
        legend=dict(orientation="h", x=0.5, xanchor="center", y=-0.14))
    fig.update_xaxes(range=[ini, x_fin], tickformat="%d %b\n%Y", gridcolor="#e6e6e6", row=1, col=1)
    fig.update_yaxes(title_text="Aportes (% de la media)", rangemode="tozero", gridcolor="#e6e6e6",
                     row=1, col=1)
    fig.update_xaxes(title_text="% de la media", gridcolor="#e6e6e6", row=2, col=1)
    fig.update_yaxes(tickmode="array", tickvals=[1, 0], ticktext=list(cat.keys()), range=[-0.6, 1.6],
                     row=2, col=1)

    datos = {
        "grupos": grupos,
        "stat": {"media": {"resto": r1(e_resto["media"]), "cierre": r1(e_cierre["media"])},
                 "mediana": {"resto": r1(e_resto["mediana"]), "cierre": r1(e_cierre["mediana"])},
                 "rango": {"resto": [r1(e_resto["min"]), r1(e_resto["max"])],
                           "cierre": [r1(e_cierre["min"]), r1(e_cierre["max"])]}},
    }
    return fig, datos


PLANTILLA = r"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITULO_PAGINA__</title>
<style>
:root{--txt:#1d2733;--sub:#5b6773;--bg:#fff;--card:#f6f7f9;--brd:#d9dde3;--azul:#1f5f8b}
@media (prefers-color-scheme: dark){:root{--txt:#e8ecf1;--sub:#a3adb8;--bg:#161b22;--card:#1f262f;--brd:#34404d}}
body{margin:0;font-family:Segoe UI,Arial,sans-serif;color:var(--txt);background:var(--bg)}
header{padding:12px 20px;border-bottom:1px solid var(--brd);text-align:center}
h1{font-size:20px;margin:0;font-weight:600} .sub{color:var(--sub);font-size:13px;margin-top:4px}
#main{display:flex;gap:16px;padding:12px 20px 20px} #izq{flex:1 1 auto;min-width:0} #der{flex:0 0 280px;display:flex;flex-direction:column;gap:12px}
.card{background:var(--card);border:1px solid var(--brd);border-radius:8px;padding:10px 12px;font-size:13px;color:var(--sub)}
.card .v{font-size:22px;color:var(--txt);margin-top:4px;font-weight:600}
.btns{display:flex;gap:6px} .btns button{flex:1;padding:8px;border:1px solid var(--brd);background:var(--bg);color:var(--txt);border-radius:6px;cursor:pointer;font-size:14px}
.btns button.on{background:var(--azul);color:#fff;border-color:var(--azul)}
table{border-collapse:collapse;font-size:12px;width:100%;margin-top:6px} th,td{border:1px solid var(--brd);padding:4px 6px;text-align:center} th{background:var(--card)}
h3{font-size:14px;margin:16px 0 0}
</style><script>__PLOTLYJS__</script></head><body>
<header><h1>__TITULO__</h1><div class="sub">__SUBTITULO__</div></header>
<div id="main"><div id="izq"><div id="grafico">__GRAFICO__</div>
<h3>Votos considerados en esta ronda</h3>__TABLA__
<h3>Trazabilidad por rondas (media de los votos vigentes en cada ronda)</h3>__HISTORIAL__</div>
<aside id="der"><div class="card">Estadístico<div class="btns" style="margin-top:6px">
<button data-s="media" class="on">Media</button><button data-s="mediana">Mediana</button><button data-s="rango">Rango</button></div></div>
<div class="card">Esperado resto del mes<div class="v" id="v-resto"></div></div>
<div class="card">Cierre de mes<div class="v" id="v-cierre"></div></div></aside></div>
<script>
const D = __DATOS__;
const gd = document.getElementById('grafico').getElementsByClassName('plotly-graph-div')[0];
const f = v => v===null?'—':(Math.round(v*10)/10)+' %';
function mostrar(s){
  document.querySelectorAll('.btns button').forEach(b => b.classList.toggle('on', b.dataset.s===s));
  const todos = [].concat(...Object.values(D.grupos));
  Plotly.restyle(gd, {visible: false}, todos);
  Plotly.restyle(gd, {visible: true}, D.grupos[s]);
  const t = D.stat[s];
  document.getElementById('v-resto').textContent = s==='rango' ? f(t.resto[0])+' – '+f(t.resto[1]) : f(t.resto);
  document.getElementById('v-cierre').textContent = s==='rango' ? f(t.cierre[0])+' – '+f(t.cierre[1]) : f(t.cierre);
}
document.querySelectorAll('.btns button').forEach(b => b.addEventListener('click', () => mostrar(b.dataset.s)));
window.addEventListener('load', () => setTimeout(() => Plotly.Plots.resize(gd), 50));
mostrar('media');
</script></body></html>
"""


def tabla_votos(vigentes, semanas, evaluadores) -> str:
    cab = "".join(f"<th>S{s['n']}<br>{etiqueta_semana_corta(s)}</th>" for s in semanas)
    filas = []
    for v in vigentes:
        celdas = []
        for s in semanas:
            val = v["semanas"].get("S%d" % s["n"])
            celdas.append("<td>%s</td>" % ("" if val is None else r1(val)))
        c = "".join(celdas)
        filas.append(f"<tr><td>{html.escape(evaluadores.get(v['evaluador'], v['evaluador']))} "
                     f"({v['evaluador']})</td><td>{v['ronda']}</td>{c}<td>{r1(v['resto'])}</td>"
                     f"<td>{r1(v['cierre'])}</td><td>{v['fecha_hora']:%Y-%m-%d %H:%M}</td></tr>")
    return (f"<table><tr><th>Evaluador</th><th>Ronda del voto</th>{cab}<th>Resto del mes</th>"
            f"<th>Cierre</th><th>Registrado</th></tr>{''.join(filas)}</table>")


def tabla_historial(historial) -> str:
    filas = "".join(f"<tr><td>{r}</td><td>{n}</td><td>{r1(m)}</td><td>{r1(c)}</td></tr>"
                    for r, n, m, c in historial)
    return ("<table><tr><th>Ronda</th><th>Evaluadores</th><th>Media resto del mes</th>"
            f"<th>Media cierre</th></tr>{filas}</table>")


def escribir_consolidacion(ruta, cuenca, ronda, vigentes, historial, serie, fechas, evaluadores, mes):
    fig, datos = construir_consolidacion(cuenca, ronda, vigentes, historial, serie, fechas, evaluadores)
    div = fig.to_html(full_html=False, include_plotlyjs=False,
                      config={"displaylogo": False, "responsive": True})
    pagina = (PLANTILLA.replace("__TITULO_PAGINA__", html.escape(f"Consolidación {cuenca} R{ronda}"))
              .replace("__TITULO__", html.escape(f"Consolidación Ronda {ronda} de votaciones · {cuenca}"))
              .replace("__SUBTITULO__", html.escape(f"Pronóstico {mes} · {len(vigentes)} evaluadores · "
                                                    f"horizonte {fechas['hor_ini']:%d/%m/%Y} – {fechas['hor_fin']:%d/%m/%Y}"))
              .replace("__TABLA__", tabla_votos(vigentes, fechas["semanas"], evaluadores))
              .replace("__HISTORIAL__", tabla_historial(historial))
              .replace("__DATOS__", json.dumps(datos).replace("</", "<\\/"))
              .replace("__GRAFICO__", div)
              .replace("__PLOTLYJS__", plotly_js()))
    ruta.write_text(pagina, encoding="utf-8")


# --------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--aportes", default=BASE / ARCHIVO_APORTES, type=Path)
    ap.add_argument("--consenso", default=BASE / ARCHIVO_CONSENSO, type=Path)
    ap.add_argument("--salida", default=BASE, type=Path, help="carpeta donde está año/mes/fecha")
    ap.add_argument("--logs", default=BASE / "logs", type=Path)
    ap.add_argument("--fecha-corrida", default=None, help="AAAA-MM-DD (por defecto: fecha de realización del consenso)")
    ap.add_argument("--ronda", type=int, default=None, help="solo esta ronda (por defecto: todas las existentes)")
    args = ap.parse_args()

    run_id = uuid.uuid4().hex[:8]
    t0 = time.time()
    ruta_log = configurar_logging(args.logs, run_id)
    log.info("=" * 70)
    log.info("INICIO consolidar_votos run_id=%s", run_id)
    log.info("Usuario=%s host=%s SO=%s python=%s pandas=%s plotly=%s", getpass.getuser(),
             platform.node(), platform.platform(), platform.python_version(),
             pd.__version__, plotly.__version__)
    log.info("Argumentos: %s", {k: str(v) for k, v in vars(args).items()})
    try:
        for p in (args.aportes, args.consenso):
            if not p.exists():
                raise FileNotFoundError(f"No se encontró el archivo Excel: {p} (debe estar en la carpeta del script)")
            st = p.stat()
            log.info("Entrada: %s | %d bytes | modificado %s | sha256=%s", p.resolve(), st.st_size,
                     datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"), sha256(p))

        fechas = leer_fechas(args.consenso)
        evaluadores = leer_evaluadores(args.consenso)
        mes = f"{fechas['hor_ini']:%Y-%m}"
        fecha_corrida = (pd.Timestamp(args.fecha_corrida).normalize() if args.fecha_corrida
                         else fechas["consenso"])
        carpeta = carpeta_corrida(args.salida, fecha_corrida)
        carpeta_votos = carpeta / "votos"
        if not carpeta_votos.exists():
            raise FileNotFoundError(f"No existe la carpeta de votos: {carpeta_votos} "
                                    "(¿se corrió crear_graficos_consenso.py para esta fecha?)")
        (carpeta / "consolidacion").mkdir(exist_ok=True)
        log.info("Carpeta de corrida: %s | mes del pronóstico: %s", carpeta.resolve(), mes)

        aportes = leer_aportes(args.aportes)
        ini, _ = definir_ventana(args.aportes, aportes)
        if ini >= fechas["hor_ini"]:
            ini = fechas["hor_ini"] - pd.Timedelta(days=30)
        cuencas = {c for c in aportes.columns if c != "Fecha"}

        votos, rechazados = leer_votos(carpeta_votos, evaluadores, cuencas, mes)
        wb, ws, cols, nuevos = escribir_votos(args.consenso, votos, args.logs)
        todos = votos_del_mes(ws, cols, mes)
        log.info("Votos del mes %s en Registros: %d", mes, len(todos))

        generados = []
        for cuenca in sorted({v["cuenca"] for v in todos}):
            vc = [v for v in todos if v["cuenca"] == cuenca]
            rondas = sorted({v["ronda"] for v in vc})
            objetivo = [args.ronda] if args.ronda else rondas
            serie = aportes.loc[(aportes["Fecha"] >= ini) & (aportes["Fecha"] < fechas["hor_ini"]),
                                ["Fecha", cuenca]].dropna() if cuenca in cuencas else aportes.iloc[0:0]
            historial = []
            for r in rondas:
                vig = votos_vigentes(vc, r)
                historial.append((r, len(vig), estadisticos([v["resto"] for v in vig])["media"],
                                  estadisticos([v["cierre"] for v in vig])["media"]))
            for r in objetivo:
                vig = votos_vigentes(vc, r)
                if not vig:
                    log.warning("Cuenca '%s': ronda %d sin votos vigentes; se omite", cuenca, r)
                    continue
                ruta = carpeta / "consolidacion" / f"Consolidacion_{nombre_archivo(cuenca)}_Ronda{r}.html"
                escribir_consolidacion(ruta, cuenca, r, vig, [h for h in historial if h[0] <= r],
                                       serie, fechas, evaluadores, mes)
                generados.append(ruta.name)
                log.info("Consolidación %s ronda %d -> %s | evaluadores=%s | media resto=%.1f cierre=%.1f",
                         cuenca, r, ruta.name, [f"{v['evaluador']}(R{v['ronda']})" for v in vig],
                         estadisticos([v["resto"] for v in vig])["media"],
                         estadisticos([v["cierre"] for v in vig])["media"])

        log.info("RESUMEN: votos leídos=%d rechazados=%d nuevos en Registros=%d | gráficos=%d",
                 len(votos), len(rechazados), len(nuevos), len(generados))
        log.info("FIN OK run_id=%s duración=%.2fs log=%s", run_id, time.time() - t0, ruta_log.resolve())
        print(f"\nVotos nuevos registrados: {len(nuevos)} | rechazados: {len(rechazados)} | "
              f"gráficos: {len(generados)}\nCarpeta: {(carpeta / 'consolidacion').resolve()}")
        return 0
    except Exception:
        log.exception("FIN CON ERROR run_id=%s tras %.2fs", run_id, time.time() - t0)
        return 1


if __name__ == "__main__":
    sys.exit(main())

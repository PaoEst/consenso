# -*- coding: utf-8 -*-
"""Lectura de Consenso.xlsx y utilidades compartidas por los scripts de consenso."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
import plotly.offline

from graficar_aportes import BASE, SALIDAS, log, norm

ARCHIVO_CONSENSO = "Consenso.xlsx"
ARCHIVO_PRONOSTICO = "Pronostico_Drivers.xlsx"
HOJA_FECHAS = "Fechas"
HOJA_EVALUADORES = "Evaluadores_Tecnicos"
HOJA_REGISTROS = "Registros"
N_SEMANAS = 5

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]


def fecha_corta(f: pd.Timestamp) -> str:
    return f"{f.day:02d} {MESES[f.month - 1][:3]}"


def etiqueta_semana_corta(s: dict) -> str:
    """'01-07 nov' o '28 oct - 03 nov' para los rótulos de semana."""
    if s["ini"].month == s["fin"].month:
        return f"{s['ini'].day:02d}-{fecha_corta(s['fin'])}"
    return f"{fecha_corta(s['ini'])} - {fecha_corta(s['fin'])}"


def _fecha(v):
    f = pd.to_datetime(v, errors="coerce", dayfirst=isinstance(v, str))
    return None if pd.isna(f) else f.normalize()


_TERMINO = re.compile(r"([+-]?)([A-Za-z]{1,2}[0-9]+|[0-9]+(?:\.[0-9]+)?)")


def _valor_celda(ws_f, ws_v, coord: str, profundidad: int = 0):
    """Valor de una celda; si no tiene valor guardado (Excel no recalculado, p. ej. tras guardar
    con openpyxl) evalúa fórmulas simples del tipo =+C2, =+E2+6, =F2-1."""
    v = ws_v[coord].value
    if v is not None:
        return v
    f = ws_f[coord].value
    if not (isinstance(f, str) and f.startswith("=")) or profundidad > 20:
        return f
    expr = f[1:].replace(" ", "")
    total, pos, fecha = 0.0, 0, False
    while pos < len(expr):
        m = _TERMINO.match(expr, pos)
        if not m:
            raise ValueError(f"Fórmula no soportada en {coord}: {f}")
        signo = -1 if m.group(1) == "-" else 1
        tok = m.group(2)
        if tok[0].isalpha():
            ref = _valor_celda(ws_f, ws_v, tok.upper(), profundidad + 1)
            if isinstance(ref, (datetime, date)):
                fecha = True
                ref = (ref.date() if isinstance(ref, datetime) else ref).toordinal()
            if ref is None:
                return None
            total += signo * float(ref)
        else:
            total += signo * float(tok)
        pos = m.end()
    return datetime.fromordinal(int(round(total))) if fecha else total


def leer_fechas(ruta: Path) -> dict:
    """Hoja Fechas (fila 2): consenso, horizonte y semanas S1..S5 (inicio/fin)."""
    wb_f = openpyxl.load_workbook(ruta)
    wb_v = openpyxl.load_workbook(ruta, data_only=True)
    if HOJA_FECHAS not in wb_f.sheetnames:
        raise ValueError(f"{ruta.name} no tiene la hoja '{HOJA_FECHAS}'")
    ws_f, ws_v = wb_f[HOJA_FECHAS], wb_v[HOJA_FECHAS]
    v = [_valor_celda(ws_f, ws_v, f"{get_column_letter(c)}2") for c in range(1, 5 + 2 * N_SEMANAS)]
    out = {"seguimiento_ini": _fecha(v[0]), "consenso": _fecha(v[1]),
           "hor_ini": _fecha(v[2]), "hor_fin": _fecha(v[3]), "semanas": []}
    for faltante in ("consenso", "hor_ini", "hor_fin"):
        if out[faltante] is None:
            raise ValueError(f"Hoja '{HOJA_FECHAS}': falta o es ilegible la fecha '{faltante}'")
    if out["hor_ini"] > out["hor_fin"]:
        raise ValueError("Hoja Fechas: el inicio del horizonte es posterior a su fin")
    for k in range(N_SEMANAS):
        ini, fin = _fecha(v[4 + 2 * k]), _fecha(v[5 + 2 * k])
        if ini is None or fin is None:
            log.info("Fechas: S%d sin fechas completas; no se usa", k + 1)
            continue
        if ini > fin:
            raise ValueError(f"Hoja Fechas: S{k + 1} tiene inicio posterior al fin")
        if ini < out["hor_ini"] or fin > out["hor_fin"]:
            log.warning("Fechas: S%d (%s -> %s) cae fuera del horizonte (%s -> %s)", k + 1,
                        ini.date(), fin.date(), out["hor_ini"].date(), out["hor_fin"].date())
        out["semanas"].append({"n": k + 1, "ini": ini, "fin": fin, "dias": (fin - ini).days + 1})
    if not out["semanas"]:
        raise ValueError("Hoja Fechas: no hay ninguna semana definida")
    log.info("Fechas: consenso=%s horizonte=%s -> %s semanas=%s", out["consenso"].date(),
             out["hor_ini"].date(), out["hor_fin"].date(),
             [(s["n"], str(s["ini"].date()), str(s["fin"].date()), s["dias"]) for s in out["semanas"]])
    return out


def leer_evaluadores(ruta: Path) -> dict[str, str]:
    """{INICIALES_MAYUSCULA: nombre} de la hoja Evaluadores_Tecnicos."""
    df = pd.read_excel(ruta, sheet_name=HOJA_EVALUADORES)
    df = df.iloc[:, :2].dropna(how="all")
    df.columns = ["nombre", "iniciales"]
    out = {}
    for _, r in df.dropna(subset=["iniciales"]).iterrows():
        ini = str(r["iniciales"]).strip().upper()
        if ini in out:
            log.warning("Evaluadores: iniciales repetidas '%s'; se conserva la primera", ini)
            continue
        out[ini] = str(r["nombre"]).strip()
    log.info("Evaluadores técnicos (%d): %s", len(out), out)
    return out


def carpeta_corrida(base: Path, fecha: pd.Timestamp) -> Path:
    """<base>/<año>/<mes>/<AAAA-MM-DD>"""
    return base / f"{fecha.year}" / f"{fecha.month:02d}" / f"{fecha:%Y-%m-%d}"


def plotly_js() -> str:
    """Código de plotly.js para incrustarlo en el HTML (funciona sin internet)."""
    return plotly.offline.get_plotlyjs()


# ==========================================================================
# Drivers climáticos (seguimiento + pronóstico) compartidos por los dos gráficos
# ==========================================================================
import html  # noqa: E402
import json  # noqa: E402
import textwrap  # noqa: E402

import plotly.graph_objects as go  # noqa: E402

from graficar_aportes import (ARCHIVO_APORTES, ARCHIVO_DRIVERS, AZUL, NEUTRO, OCRE,  # noqa: E402,F401
                              color_impacto, drivers_de_cuenca, leer_drivers, leer_prioridad,
                              mezcla, segmentos, texto_comentarios)

ARCHIVO_INDICE = "indice_cuencas.json"   # archivo -> nombre de cuenca (para el índice/selector)
CARPETA_CONSOLIDADO = "consolidado"
PREFIJO_CONSOLIDADO = "Consolidado"


def hoja_prioridad(ruta: Path) -> str:
    """Nombre de la hoja de prioridad (tolera 'Prioridad' y el typo 'Priodidad')."""
    for h in pd.ExcelFile(ruta).sheet_names:
        if norm(h).startswith("prio"):
            return h
    raise ValueError(f"{ruta.name}: no se encontró una hoja de prioridad")


def cargar_drivers(ruta_seg: Path, ruta_pro: Path) -> dict:
    """Lee seguimiento y pronóstico y define el orden de las filas del panel de drivers.

    Orden: prioridad de Seguimiento_Drivers; después las variables que solo existen en el
    pronóstico (prioridad de Pronostico_Drivers); al final cualquier otra no listada."""
    pri_seg = leer_prioridad(ruta_seg)
    drv_seg = leer_drivers(ruta_seg, pri_seg)
    pri_pro = leer_prioridad(ruta_pro, hoja_prioridad(ruta_pro))
    drv_pro = leer_drivers(ruta_pro, pri_pro)
    orden = list(pri_seg["Variable"])
    conocidas = set(pri_seg["clave"])
    for k, v in zip(pri_pro["clave"], pri_pro["Variable"]):
        if k not in conocidas:
            orden.append(v)
            conocidas.add(k)
    for df in (drv_seg, drv_pro):
        for k, v in df.drop_duplicates("clave")[["clave", "Variable"]].values:
            if k and k not in conocidas:
                orden.append(str(v).strip())
                conocidas.add(k)
    nombre_de = {}
    for v in orden:
        nombre_de.setdefault(norm(v), v)
    log.info("Variables solo del pronóstico (al final): %s",
             [v for v in orden if norm(v) not in set(pri_seg["clave"])])
    return {"pri_seg": pri_seg, "drv_seg": drv_seg, "pri_pro": pri_pro, "drv_pro": drv_pro,
            "orden": orden, "nombre_de": nombre_de}


def drivers_combinados(cuenca, ctx: dict, ini, fin):
    """Barras de seguimiento y pronóstico de una cuenca + lista ordenada de variables con barras."""
    sel_seg = drivers_de_cuenca(ctx["drv_seg"], ctx["pri_seg"], cuenca, ini, fin)
    sel_pro = drivers_de_cuenca(ctx["drv_pro"], ctx["pri_pro"], cuenca, ini, fin)
    if len(sel_seg) and len(sel_pro):  # si un día de una variable está en ambos, gana el seguimiento
        ya = set(zip(sel_seg["clave"], sel_seg["Fecha"]))
        dup = [(k, f) in ya for k, f in zip(sel_pro["clave"], sel_pro["Fecha"])]
        if any(dup):
            log.warning("%s: %d filas de pronóstico coinciden con días de seguimiento; se omiten",
                        cuenca, sum(dup))
        sel_pro = sel_pro[[not d for d in dup]]
    partes = []
    for sel, origen in ((sel_seg, "Seguimiento"), (sel_pro, "Pronóstico")):
        s = segmentos(sel)
        s["origen"] = origen
        partes.append(s)
    seg = pd.concat(partes, ignore_index=True)
    if len(seg):
        seg["Variable"] = seg["clave"].map(lambda k: ctx["nombre_de"].get(k, k))
    presentes = set(seg["clave"])
    return seg, [v for v in ctx["orden"] if norm(v) in presentes]


def real_a_la_fecha(aportes: pd.DataFrame, cuenca: str, hor_ini: pd.Timestamp) -> list[float]:
    """Aportes con dato del mes del inicio del horizonte, anteriores al horizonte (no mezcla meses)."""
    f = aportes["Fecha"]
    m = aportes[(f.dt.year == hor_ini.year) & (f.dt.month == hor_ini.month) & (f < hor_ini)][cuenca]
    return [float(v) for v in m.dropna()]


def hover_barra(r) -> str:
    imp = "s/d" if pd.isna(r["impacto"]) else f"{r['impacto']:g}"
    dias = (r["fin"] - r["ini"]).days + 1
    return (f"<b>{html.escape(r['Variable'])}</b> ({r['origen']})<br>"
            f"{r['ini']:%d %b} – {r['fin']:%d %b %Y} ({dias} d)<br>"
            f"Nota: {html.escape(r['nota']) or '—'}<br>"
            f"Impacto estimado: {imp}<br>Lugar: {html.escape(str(r['lugar']))}"
            + texto_comentarios(r["coment"], r["ini"], r["fin"]))


def etiqueta_y(v: str, ancho: int = 24) -> str:
    """Parte los nombres largos de variable en dos renglones."""
    return "<br>".join(textwrap.wrap(v, ancho)[:2]) if len(v) > ancho else v


def altura_panel_drivers(n_vars: int) -> int:
    return 26 * max(n_vars, 1) + 50


def agregar_barras_drivers(fig, seg: pd.DataFrame, orden_vars: list[str], row: int, col: int = 1):
    """Dibuja las barras de drivers (una fila por variable) y configura el eje Y de esa fila."""
    pos = {v: i for i, v in enumerate(orden_vars)}
    h = 0.36
    cx, cy, ct = [], [], []
    for _, r in seg.iterrows():
        yc = pos[r["Variable"]]
        fin_b = r["fin"] + pd.Timedelta(days=1)
        fig.add_trace(go.Scatter(
            x=[r["ini"], fin_b, fin_b, r["ini"], r["ini"]],
            y=[yc - h, yc - h, yc + h, yc + h, yc - h], mode="lines", fill="toself",
            fillcolor=color_impacto(r["impacto"])[0], line=dict(color="white", width=1.5),
            hoveron="fills", text=hover_barra(r), hoverinfo="text", showlegend=False),
            row=row, col=col)
        if r["nota"]:
            cx.append(r["ini"] + (fin_b - r["ini"]) / 2)
            cy.append(yc)
            ct.append(r["nota"])
    if ct:
        fig.add_trace(go.Scatter(x=cx, y=cy, mode="text", text=ct, hoverinfo="skip",
                                 textfont=dict(color="white", size=11), showlegend=False),
                      row=row, col=col)
    for nombre, color in (("Favorece humedad", mezcla(*AZUL, 0.8)),
                          ("Favorece condiciones secas", mezcla(*OCRE, 0.8)),
                          ("Presente, efecto neutro/incierto", NEUTRO)):
        fig.add_trace(go.Scatter(x=[None], y=[None], mode="markers", name=nombre,
                                 marker=dict(color=color, size=12, symbol="square")),
                      row=row, col=col)
    fig.update_yaxes(tickmode="array", tickvals=list(range(len(orden_vars))),
                     ticktext=[etiqueta_y(v) for v in orden_vars],
                     range=[len(orden_vars) - 0.4, -0.6], row=row, col=col, gridcolor="#efefef",
                     tickfont=dict(size=10), automargin=True)


# ==========================================================================
# Navegación entre cuencas (índice + selector) con enlaces relativos
# ==========================================================================
def registrar_cuenca(carpeta: Path, archivo: str, cuenca: str) -> None:
    ruta = carpeta / ARCHIVO_INDICE
    d = json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}
    d[archivo] = cuenca
    ruta.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")


def cuencas_registradas(carpeta: Path) -> dict[str, str]:
    """{archivo_base: nombre} (archivo_base = nombre_archivo(cuenca))."""
    ruta = carpeta / ARCHIVO_INDICE
    return json.loads(ruta.read_text(encoding="utf-8")) if ruta.exists() else {}


def barra_navegacion(carpeta: Path, actual: str, donde: str, ronda: int | None = None) -> str:
    """HTML de la barra superior: enlace al índice y selector de cuenca (y de ronda).

    donde = 'votacion' (página en la carpeta de corrida) | 'consolidado' (en carpeta/consolidado/)."""
    cuencas = cuencas_registradas(carpeta)
    en_cons = donde == "consolidado"
    ruta_cons = carpeta / CARPETA_CONSOLIDADO
    prefijo = "../" if en_cons else ""
    opciones = []
    for archivo, nombre in cuencas.items():
        if en_cons:
            if ronda is None:
                continue
            # misma ronda si existe; si no, la más cercana anterior (o la primera disponible)
            disp = sorted(int(m.group(1)) for f in ruta_cons.glob(f"{PREFIJO_CONSOLIDADO}_{archivo}_Ronda*.html")
                          if (m := re.search(r"Ronda(\d+)\.html$", f.name)))
            if not disp:
                continue
            r_dest = ronda if ronda in disp else max([r for r in disp if r < ronda] or [disp[0]])
            destino = f"{PREFIJO_CONSOLIDADO}_{archivo}_Ronda{r_dest}.html"
        else:
            destino = f"Consenso_{archivo}.html"
            if not (carpeta / destino).exists():
                continue
        sel = " selected" if archivo == actual else ""
        opciones.append(f'<option value="{html.escape(destino)}"{sel}>{html.escape(nombre)}</option>')
    partes = [f'<a href="{prefijo}index.html">‹ Índice</a>']
    if opciones:
        partes.append('<label>Cuenca <select onchange="location.href=this.value">'
                      + "".join(opciones) + "</select></label>")
    if en_cons and ronda is not None:
        rondas = sorted({int(m.group(1)) for f in ruta_cons.glob(f"{PREFIJO_CONSOLIDADO}_{actual}_Ronda*.html")
                         if (m := re.search(r"Ronda(\d+)\.html$", f.name))})
        if len(rondas) > 1:
            ops = "".join(f'<option value="{PREFIJO_CONSOLIDADO}_{actual}_Ronda{r}.html"'
                          f'{" selected" if r == ronda else ""}>Ronda {r}</option>' for r in rondas)
            partes.append(f'<label>Ronda <select onchange="location.href=this.value">{ops}</select></label>')
    return '<nav class="nav">' + "".join(partes) + "</nav>"


ESTILO_NAV = ("<style>.nav{display:flex;flex-wrap:wrap;gap:18px;align-items:center;padding:8px 20px;"
              "font-size:13px;border-bottom:1px solid var(--brd)} .nav a{color:var(--azul);text-decoration:none;"
              "font-weight:600} .nav select{font-size:13px;padding:3px 6px;margin-left:6px;border-radius:5px;"
              "border:1px solid var(--brd);background:var(--bg);color:var(--txt)}</style>")


def escribir_indice(carpeta: Path, fechas: dict) -> Path:
    """index.html de la corrida: una fila por cuenca con enlace a votación y a cada ronda del Consolidado."""
    cuencas = cuencas_registradas(carpeta)
    filas = []
    for archivo, nombre in cuencas.items():
        if not (carpeta / f"Consenso_{archivo}.html").exists():
            continue
        enlaces = [f'<a href="Consenso_{html.escape(archivo)}.html">Votación</a>']
        for f in sorted((carpeta / CARPETA_CONSOLIDADO).glob(f"{PREFIJO_CONSOLIDADO}_{archivo}_Ronda*.html")):
            m = re.search(r"Ronda(\d+)\.html$", f.name)
            enlaces.append(f'<a href="{CARPETA_CONSOLIDADO}/{html.escape(f.name)}">Consolidado ronda {m.group(1)}</a>')
        filas.append(f"<tr><td>{html.escape(nombre)}</td><td>{' · '.join(enlaces)}</td></tr>")
    pagina = f"""<!doctype html><html lang="es"><head><meta charset="utf-8">
<title>Consenso de aportes</title><style>
:root{{--txt:#1d2733;--sub:#5b6773;--bg:#fff;--brd:#d9dde3;--azul:#1f5f8b}}
@media (prefers-color-scheme: dark){{:root{{--txt:#e8ecf1;--sub:#a3adb8;--bg:#161b22;--brd:#34404d;--azul:#7fb3e6}}}}
body{{font-family:Segoe UI,Arial,sans-serif;color:var(--txt);background:var(--bg);margin:0;padding:28px 36px}}
h1{{font-size:22px;margin:0 0 4px}} p{{color:var(--sub);margin:0 0 18px;font-size:14px}}
table{{border-collapse:collapse;min-width:420px}} td,th{{border-bottom:1px solid var(--brd);padding:9px 16px;text-align:left;font-size:14px}}
th{{color:var(--sub);font-weight:600}} a{{color:var(--azul);font-weight:600;text-decoration:none}} a:hover{{text-decoration:underline}}
</style></head><body><h1>Consenso de aportes por cuenca</h1>
<p>Pronóstico {MESES[fechas['hor_ini'].month - 1]} {fechas['hor_ini'].year} · horizonte {fechas['hor_ini']:%d/%m/%Y} – {fechas['hor_fin']:%d/%m/%Y}
· consenso del {fechas['consenso']:%d/%m/%Y}</p>
<table><tr><th>Cuenca</th><th>Abrir</th></tr>{''.join(filas)}</table></body></html>"""
    ruta = carpeta / "index.html"
    ruta.write_text(pagina, encoding="utf-8")
    return ruta

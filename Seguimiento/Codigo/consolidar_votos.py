#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Registra los votos de los evaluadores en Consenso.xlsx > Registros, escribe el resultado del
consenso en la hoja Consolidado y genera, en la misma corrida, el gráfico Consolidado por
cuenca y por ronda.

Flujo
  1. Cada evaluador abre su Consenso_<cuenca>.html, vota y pulsa GUARDAR (se descarga un
     archivo voto_*.json). Los archivos se dejan en  <año>/<mes>/<fecha>/votos/
  2. Se corre este script. Cada voto nuevo se agrega a Registros (sin duplicar); por cada cuenca y
     ronda se escriben 3 filas (Media, Mediana, Rango) en la hoja Consolidado y se crea
     <año>/<mes>/<fecha>/consolidado/Consolidado_<cuenca>_Ronda<N>.html

Rondas: el "No Evaluación" es la ronda. En la ronda N cada evaluador cuenta con su voto más
reciente de ronda <= N (si no volvió a votar, se mantiene el de la ronda anterior).

Uso:  python consolidar_votos.py [--ronda N] [--fecha-corrida AAAA-MM-DD]
"""
from __future__ import annotations

import argparse
import getpass
import html
import json
import os
import platform
import shutil
import statistics
import sys
import time
import uuid
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd
import plotly
import plotly.graph_objects as go
from plotly.subplots import make_subplots

from consenso_comun import (ARCHIVO_CONSENSO, ARCHIVO_PRONOSTICO, BASE, CARPETA_CONSOLIDADO,
                            ESTILO_NAV, HOJA_REGISTROS, N_SEMANAS, PREFIJO_CONSOLIDADO, SALIDAS,
                            agregar_barras_drivers, altura_panel_drivers, barra_navegacion,
                            cargar_drivers, carpeta_corrida, cuencas_registradas, drivers_combinados,
                            escribir_indice, etiqueta_semana_corta, leer_evaluadores, leer_fechas,
                            plotly_js, real_a_la_fecha, registrar_cuenca)
from graficar_aportes import (ARCHIVO_APORTES, ARCHIVO_DRIVERS, C_ABAJO, C_ARRIBA, C_LINEA,
                              configurar_logging, definir_ventana, leer_aportes, log,
                              nombre_archivo, norm, relleno_media, sha256)

COLUMNAS_BASE = ["Fecha_hora_registro", "No Evaluación", "Cuenca", "Evaluador", "Aportes S1",
                 "Aportes S2", "Aportes S3", "Aportes S4", "Aportes S5", "Real a la fecha",
                 "Aportes_resto del mes", "Aportes cierre de mes"]
COLUMNAS_EXTRA = ["Mes_pronostico", "Id_voto"]
HOJA_CONSOLIDADO = "Consolidado"
COLUMNAS_CONSOLIDADO = COLUMNAS_BASE + ["Estadístico", "Mes_pronostico"]
ESTADISTICOS = ("Media", "Mediana", "Rango")
PALETA = ["#1f77b4", "#ff7f0e", "#2ca02c", "#9467bd", "#8c564b", "#e377c2", "#17becf", "#bcbd22"]


# --------------------------------------------------------------------------
# Lectura y validación de votos
# --------------------------------------------------------------------------
def _num(v):
    if v is None or v == "":
        return None
    f = float(v)
    if f != f:  # NaN
        raise ValueError("valor no numérico")
    return f


def leer_votos(carpeta_votos: Path, evaluadores: dict, cuencas: set[str], mes: str):
    """Devuelve (votos_validos, rechazados). Cada voto es un dict normalizado."""
    validos, rechazados, ids = [], [], set()
    archivos = sorted(carpeta_votos.glob("*.json"))
    log.info("Votos: %d archivos .json en %s", len(archivos), carpeta_votos)
    for f in archivos:
        try:
            d = json.loads(f.read_text(encoding="utf-8-sig"))
            vid = str(d["id"]).strip()
            ev = str(d["evaluador"]).strip().upper()
            if ev not in evaluadores:
                raise ValueError(f"evaluador '{ev}' no está en Evaluadores_Tecnicos")
            ronda = int(d["no_evaluacion"])
            if ronda < 1 or ronda != float(d["no_evaluacion"]):
                raise ValueError("No Evaluación debe ser un entero >= 1")
            cuenca = str(d["cuenca"]).strip()
            if cuenca not in cuencas:
                raise ValueError(f"cuenca '{cuenca}' no existe en Seguimiento_Aportes")
            if str(d["mes_pronostico"]) != mes:
                raise ValueError(f"el voto es del mes {d['mes_pronostico']}, la corrida es de {mes}")
            sem = {f"S{k}": _num(d["semanas"].get(f"S{k}")) for k in range(1, N_SEMANAS + 1)}
            if any(v is not None and v < 0 for v in sem.values()):
                raise ValueError("hay aportes negativos")
            if all(v is None for v in sem.values()):
                raise ValueError("el voto no tiene ninguna semana")
            resto, cierre = _num(d["resto_del_mes"]), _num(d["cierre_de_mes"])
            if resto is None or cierre is None:
                raise ValueError("faltan resto del mes o cierre de mes")
            if vid in ids:
                raise ValueError(f"id de voto repetido en otro archivo ({vid})")
            ids.add(vid)
            validos.append({
                "id": vid, "fecha_hora": datetime.fromisoformat(str(d["fecha_hora"])),
                "evaluador": ev, "ronda": ronda, "cuenca": cuenca, "semanas": sem,
                "real": _num(d.get("real_a_la_fecha")), "resto": resto, "cierre": cierre,
                "mes": mes, "archivo": f.name})
            log.debug("Voto válido: %s | %s | %s | ronda %d | %s", f.name, ev, cuenca, ronda, sem)
        except Exception as e:
            rechazados.append((f.name, str(e)))
            log.warning("Voto RECHAZADO %s: %s", f.name, e)
    log.info("Votos válidos=%d rechazados=%d", len(validos), len(rechazados))
    return validos, rechazados


# --------------------------------------------------------------------------
# Registros (Excel)
# --------------------------------------------------------------------------
def cargar_registros(ruta: Path):
    wb = openpyxl.load_workbook(ruta)
    if HOJA_REGISTROS not in wb.sheetnames:
        raise ValueError(f"{ruta.name} no tiene la hoja '{HOJA_REGISTROS}'")
    ws = wb[HOJA_REGISTROS]
    cols = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if v is not None:
            cols[norm(v)] = c
    for extra in COLUMNAS_EXTRA:
        if norm(extra) not in cols:
            c = max(cols.values()) + 1
            ws.cell(1, c, extra)
            cols[norm(extra)] = c
            log.info("Registros: se agregó la columna '%s' en la posición %d", extra, c)
    faltan = [c for c in COLUMNAS_BASE if norm(c) not in cols]
    if faltan:
        raise ValueError(f"Registros: faltan las columnas {faltan}")
    return wb, ws, cols


def ultima_fila(ws, cols) -> int:
    """Última fila con algún dato en las columnas de Registros."""
    for r in range(ws.max_row, 1, -1):
        if any(ws.cell(r, c).value not in (None, "") for c in cols.values()):
            return r
    return 1


def agregar_votos(ws, cols, votos: list[dict]) -> list[dict]:
    """Agrega en memoria (hoja Registros) los votos que aún no están. Devuelve los nuevos."""
    col_id = cols[norm("Id_voto")]
    ya = {str(ws.cell(r, col_id).value) for r in range(2, ws.max_row + 1) if ws.cell(r, col_id).value}
    nuevos = [v for v in votos if v["id"] not in ya]
    log.info("Registros: %d votos ya estaban registrados, %d son nuevos", len(votos) - len(nuevos), len(nuevos))
    fila = ultima_fila(ws, cols) + 1
    for v in sorted(nuevos, key=lambda x: (x["fecha_hora"], x["id"])):
        vals = {"Fecha_hora_registro": v["fecha_hora"], "No Evaluación": v["ronda"],
                "Cuenca": v["cuenca"], "Evaluador": v["evaluador"],
                "Real a la fecha": v["real"], "Aportes_resto del mes": v["resto"],
                "Aportes cierre de mes": v["cierre"], "Mes_pronostico": v["mes"],
                "Id_voto": v["id"]}
        for k in range(1, N_SEMANAS + 1):
            vals[f"Aportes S{k}"] = v["semanas"][f"S{k}"]
        for nombre, val in vals.items():
            ws.cell(fila, cols[norm(nombre)], val)
        ws.cell(fila, cols[norm("Fecha_hora_registro")]).number_format = "yyyy-mm-dd hh:mm:ss"
        log.info("Registros fila %d <- %s | %s | ronda %d | resto=%s cierre=%s | id=%s", fila,
                 v["evaluador"], v["cuenca"], v["ronda"], v["resto"], v["cierre"], v["id"])
        fila += 1
    return nuevos


def guardar_excel(wb, ruta: Path, dir_logs: Path) -> None:
    """Copia de seguridad + guardado atómico del libro."""
    dir_bk = dir_logs / "backups"
    dir_bk.mkdir(parents=True, exist_ok=True)
    bk = dir_bk / f"{ruta.stem}_{datetime.now():%Y%m%d_%H%M%S}{ruta.suffix}"
    shutil.copy2(ruta, bk)
    log.info("Copia de seguridad: %s", bk.resolve())
    tmp = ruta.with_name(ruta.stem + ".tmp" + ruta.suffix)
    wb.save(tmp)
    try:
        os.replace(tmp, ruta)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise PermissionError(f"No se pudo escribir {ruta.name}: ciérralo en Excel y vuelve a correr "
                              f"(no se registró nada; la copia {bk.name} no cambió nada).")
    log.info("Guardado %s", ruta.resolve())


def votos_del_mes(ws, cols, mes: str) -> list[dict]:
    """Todos los votos del mes registrados en Registros (todas las corridas)."""
    g = lambda r, n: ws.cell(r, cols[norm(n)]).value
    out, omitidos = [], 0
    for r in range(2, ws.max_row + 1):
        cuenca, ev, ronda, m = g(r, "Cuenca"), g(r, "Evaluador"), g(r, "No Evaluación"), g(r, "Mes_pronostico")
        if not cuenca or not ev or ronda in (None, ""):
            omitidos += 1 if any(ws.cell(r, c).value not in (None, "") for c in cols.values()) else 0
            continue
        if str(m) != mes:
            continue
        fh = g(r, "Fecha_hora_registro")
        out.append({"cuenca": str(cuenca).strip(), "evaluador": str(ev).strip().upper(),
                    "ronda": int(ronda), "fecha_hora": fh if isinstance(fh, datetime) else datetime.min,
                    "semanas": {f"S{k}": g(r, f"Aportes S{k}") for k in range(1, N_SEMANAS + 1)},
                    "resto": g(r, "Aportes_resto del mes"), "cierre": g(r, "Aportes cierre de mes"),
                    "real": g(r, "Real a la fecha")})
    if omitidos:
        log.warning("Registros: %d filas incompletas o sin cuenca/mes se ignoran en la consolidación", omitidos)
    return out


# --------------------------------------------------------------------------
# Consolidado
# --------------------------------------------------------------------------
def votos_vigentes(votos_cuenca: list[dict], ronda: int) -> list[dict]:
    """Por evaluador, su voto más reciente con ronda <= ronda (arrastra rondas anteriores)."""
    por_ev = {}
    for v in sorted(votos_cuenca, key=lambda x: (x["ronda"], x["fecha_hora"])):
        if v["ronda"] <= ronda:
            if v["evaluador"] in por_ev and por_ev[v["evaluador"]]["ronda"] == v["ronda"]:
                log.warning("Evaluador %s votó más de una vez la ronda %d (%s); se usa el último",
                            v["evaluador"], v["ronda"], v["cuenca"])
            por_ev[v["evaluador"]] = v
    return sorted(por_ev.values(), key=lambda x: x["evaluador"])


def estadisticos(valores) -> dict:
    v = [float(x) for x in valores if x is not None]
    if not v:
        return {"media": None, "mediana": None, "min": None, "max": None}
    return {"media": sum(v) / len(v), "mediana": statistics.median(v), "min": min(v), "max": max(v)}


def r1(x):
    return None if x is None else round(float(x), 1)


def txt_rango(mn, mx) -> str:
    if mn is None:
        return ""
    return f"{round(mn, 1):g}" if round(mn, 1) == round(mx, 1) else f"{round(mn, 1):g} - {round(mx, 1):g}"


def estadisticos_ronda(vigentes: list[dict], semanas: list[dict]) -> dict:
    """Todos los estadísticos de una ronda: por semana, resto del mes y cierre."""
    return {
        "n": len(vigentes),
        "sem": {f"S{s['n']}": estadisticos([v["semanas"].get(f"S{s['n']}") for v in vigentes])
                for s in semanas},
        "resto": estadisticos([v["resto"] for v in vigentes]),
        "cierre": estadisticos([v["cierre"] for v in vigentes]),
    }


def filas_consolidado(cuenca, ronda, vigentes, est, real_mean, mes, ahora) -> list[dict]:
    """3 filas (Media, Mediana, Rango) para la hoja Consolidado."""
    evaluadores = ", ".join(v["evaluador"] for v in vigentes)
    filas = []
    for nombre, clave in (("Media", "media"), ("Mediana", "mediana"), ("Rango", "rango")):
        def valor(e):
            if clave == "rango":
                return txt_rango(e["min"], e["max"]) or None
            return None if e[clave] is None else round(e[clave], 2)
        fila = {"Fecha_hora_registro": ahora, "No Evaluación": ronda, "Cuenca": cuenca,
                "Evaluador": evaluadores, "Real a la fecha": None if real_mean is None else round(real_mean, 2),
                "Aportes_resto del mes": valor(est["resto"]),
                "Aportes cierre de mes": valor(est["cierre"]), "Estadístico": nombre,
                "Mes_pronostico": mes}
        for k in range(1, N_SEMANAS + 1):
            fila[f"Aportes S{k}"] = valor(est["sem"][f"S{k}"]) if f"S{k}" in est["sem"] else None
        filas.append(fila)
    return filas


def _iguales(a, b) -> bool:
    if a in (None, "") and b in (None, ""):
        return True
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return str(a) == str(b)


def escribir_hoja_consolidado(wb, filas: list[dict]) -> tuple[int, int]:
    """Crea/actualiza la hoja Consolidado. Clave: (Mes_pronostico, Cuenca, No Evaluación, Estadístico).
    Devuelve (filas_nuevas, filas_actualizadas); una fila idéntica no se toca."""
    if HOJA_CONSOLIDADO not in wb.sheetnames:
        ws = wb.create_sheet(HOJA_CONSOLIDADO)
        ref = wb[HOJA_REGISTROS]
        for c, nombre in enumerate(COLUMNAS_CONSOLIDADO, start=1):
            celda = ws.cell(1, c, nombre)
            origen = ref.cell(1, min(c, ref.max_column))
            if origen.has_style:
                celda._style = origen._style
            ws.column_dimensions[celda.column_letter].width = (
                34 if nombre == "Evaluador" else 18 if nombre in ("Fecha_hora_registro", "Aportes_resto del mes", "Aportes cierre de mes") else 13)
        log.info("Se creó la hoja '%s'", HOJA_CONSOLIDADO)
    ws = wb[HOJA_CONSOLIDADO]
    cols = {norm(ws.cell(1, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(1, c).value}
    for nombre in COLUMNAS_CONSOLIDADO:
        if norm(nombre) not in cols:
            c = max(cols.values()) + 1
            ws.cell(1, c, nombre)
            cols[norm(nombre)] = c
    clave = lambda r: (str(r.get("Mes_pronostico")), str(r.get("Cuenca")), str(r.get("No Evaluación")),
                       str(r.get("Estadístico")))
    existentes = {}
    for r in range(2, ws.max_row + 1):
        reg = {n: ws.cell(r, c).value for n, c in ((x, cols[norm(x)]) for x in COLUMNAS_CONSOLIDADO)}
        if reg.get("Cuenca"):
            existentes[clave(reg)] = r
    nuevas = actualizadas = 0
    sig = max(existentes.values(), default=1) + 1
    for f in filas:
        k = clave(f)
        if k in existentes:
            r = existentes[k]
            actual = {n: ws.cell(r, cols[norm(n)]).value for n in COLUMNAS_CONSOLIDADO}
            if all(_iguales(actual[n], f[n]) for n in COLUMNAS_CONSOLIDADO if n != "Fecha_hora_registro"):
                continue
            actualizadas += 1
        else:
            r = sig
            sig += 1
            nuevas += 1
        for n in COLUMNAS_CONSOLIDADO:
            ws.cell(r, cols[norm(n)], f[n])
        ws.cell(r, cols[norm("Fecha_hora_registro")]).number_format = "yyyy-mm-dd hh:mm:ss"
    return nuevas, actualizadas


C_ESTAD = "#111111"


def construir_consolidado(cuenca, vigentes, serie, seg, orden_vars, fechas, est, evaluadores):
    hor_ini, hor_fin = fechas["hor_ini"], fechas["hor_fin"]
    x_fin = hor_fin + pd.Timedelta(days=1)
    semanas = fechas["semanas"]
    ini = serie["Fecha"].min() if len(serie) else hor_ini - pd.Timedelta(days=30)
    alto_drv = altura_panel_drivers(len(orden_vars))
    alto_tot = int(430 + alto_drv + 110)
    frac = min(0.55, max(0.25, alto_drv / alto_tot))
    fig = make_subplots(rows=2, cols=1, shared_xaxes=True, vertical_spacing=0.04,
                        row_heights=[1 - frac, frac])

    if len(serie):  # seguimiento de aportes (dato diario al centro de su día)
        x = list(serie["Fecha"] + pd.Timedelta(hours=12))
        y = list(serie[cuenca])
        xs, arr, abj = relleno_media(x, y)
        base = [100.0] * len(xs)
        for techo, color in ((arr, C_ARRIBA), (abj, C_ABAJO)):
            fig.add_trace(go.Scatter(x=xs, y=base, mode="lines", line=dict(width=0),
                                     hoverinfo="skip", showlegend=False), row=1, col=1)
            fig.add_trace(go.Scatter(x=xs, y=techo, mode="lines", line=dict(width=0), fill="tonexty",
                                     fillcolor=color, hoverinfo="skip", showlegend=False), row=1, col=1)
        fig.add_trace(go.Scatter(
            x=x, y=y, mode="lines", line=dict(color=C_LINEA, width=2.2), showlegend=False,
            hovertemplate="%{x|%d %b %Y}<br><b>%{y:.1f} %</b> de la media<extra></extra>"),
            row=1, col=1)
        fig.add_hline(y=100, line=dict(color="#333", width=1, dash="dash"), row=1, col=1)

    def escalon(valores):
        xs, ys = [], []
        for s in semanas:
            val = valores.get(f"S{s['n']}")
            if val is None:
                if xs and xs[-1] is not None:
                    xs.append(None)
                    ys.append(None)
                continue
            xs += [s["ini"], s["fin"] + pd.Timedelta(days=1)]
            ys += [val, val]
        return xs, ys

    colores = {}
    for i, v in enumerate(vigentes):  # una línea por evaluador
        colores[v["evaluador"]] = PALETA[i % len(PALETA)]
        xs, ys = escalon(v["semanas"])
        nombre = evaluadores.get(v["evaluador"], v["evaluador"])
        fig.add_trace(go.Scatter(
            x=xs, y=ys, mode="lines", name=f"{v['evaluador']} (R{v['ronda']})",
            line=dict(color=colores[v["evaluador"]], width=2),
            hovertemplate=f"<b>{html.escape(nombre)}</b> · ronda {v['ronda']}<br>%{{y:.1f}} %<extra></extra>"),
            row=1, col=1)

    grupos = {"media": [], "mediana": [], "rango": []}  # índices de trazas por estadístico
    for clave, nombre, dash in (("media", "Media", "solid"), ("mediana", "Mediana", "dash")):
        xs, ys = escalon({k: e[clave] for k, e in est["sem"].items()})
        grupos[clave].append(len(fig.data))
        fig.add_trace(go.Scatter(x=xs, y=ys, mode="lines", name=f"{nombre} del consenso",
                                 line=dict(color=C_ESTAD, width=4, dash=dash), visible=(clave == "media"),
                                 hovertemplate=f"{nombre}: %{{y:.1f}} %<extra></extra>"), row=1, col=1)
    xs_min, ys_min = escalon({k: e["min"] for k, e in est["sem"].items()})
    xs_max, ys_max = escalon({k: e["max"] for k, e in est["sem"].items()})
    grupos["rango"].append(len(fig.data))
    fig.add_trace(go.Scatter(x=xs_min, y=ys_min, mode="lines", line=dict(color=C_ESTAD, width=1),
                             name="Mínimo", visible=False,
                             hovertemplate="Mínimo: %{y:.1f} %<extra></extra>"), row=1, col=1)
    grupos["rango"].append(len(fig.data))
    fig.add_trace(go.Scatter(x=xs_max, y=ys_max, mode="lines", line=dict(color=C_ESTAD, width=1),
                             fill="tonexty", fillcolor="rgba(17,17,17,0.15)", name="Máximo",
                             visible=False, hovertemplate="Máximo: %{y:.1f} %<extra></extra>"),
                  row=1, col=1)

    agregar_barras_drivers(fig, seg, orden_vars, row=2)  # drivers climáticos

    formas = [dict(type="rect", xref="x", yref="paper", x0=hor_ini, x1=x_fin, y0=0, y1=1,
                   fillcolor="rgba(31,95,139,0.05)", line=dict(width=0), layer="below"),
              dict(type="line", xref="x", yref="paper", x0=hor_ini, x1=hor_ini, y0=0, y1=1,
                   line=dict(color="#d62728", width=2, dash="dash"))]
    for s in semanas:
        if s["ini"] > hor_ini:
            formas.append(dict(type="line", xref="x", yref="paper", x0=s["ini"], x1=s["ini"],
                               y0=0, y1=1, line=dict(color="#2e9e4f", width=1, dash="dot")))
    anot = [dict(x=ini + (hor_ini - ini) / 2, y=1.0, xref="x", yref="paper", yanchor="bottom",
                 text="<b>Seguimiento</b>", showarrow=False, font=dict(size=13)),
            dict(x=hor_ini + (x_fin - hor_ini) / 2, y=1.0, xref="x", yref="paper", yanchor="bottom",
                 text="<b>Pronóstico</b>", showarrow=False, font=dict(size=13))]
    fig.update_yaxes(title_text="Aportes (% de la media)", rangemode="tozero", gridcolor="#e6e6e6",
                     row=1, col=1)
    for fila in (1, 2):
        fig.update_xaxes(range=[ini, x_fin], matches="x" if fila == 2 else None, gridcolor="#e6e6e6",
                         showspikes=True, spikemode="across", spikethickness=1, spikecolor="#888",
                         tickformat="%d %b\n%Y", row=fila, col=1)
    fig.update_layout(height=alto_tot, template="plotly_white", shapes=formas, hovermode="closest",
                      annotations=list(fig.layout.annotations) + anot, hoverlabel=dict(align="left"),
                      margin=dict(l=20, r=20, t=40, b=70), plot_bgcolor="#fbfbf9",
                      legend=dict(orientation="h", x=0.5, xanchor="center", y=-0.1))
    return fig, grupos


PLANTILLA = r"""<!doctype html>
<html lang="es"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>__TITULO_PAGINA__</title>
__ESTILO_NAV__
<style>
:root{--txt:#1d2733;--sub:#5b6773;--bg:#fff;--card:#f6f7f9;--brd:#d9dde3;--azul:#1f5f8b}
@media (prefers-color-scheme: dark){:root{--txt:#e8ecf1;--sub:#a3adb8;--bg:#161b22;--card:#1f262f;--brd:#34404d;--azul:#7fb3e6}}
body{margin:0;font-family:Segoe UI,Arial,sans-serif;color:var(--txt);background:var(--bg)}
header{padding:12px 20px;border-bottom:1px solid var(--brd);text-align:center}
h1{font-size:20px;margin:0;font-weight:600} .sub{color:var(--sub);font-size:13px;margin-top:4px}
#main{display:flex;gap:16px;padding:12px 20px 20px} #izq{flex:1 1 auto;min-width:0} #der{flex:0 0 250px;display:flex;flex-direction:column;gap:12px}
.card{background:var(--card);border:1px solid var(--brd);border-radius:8px;padding:10px 12px;font-size:13px;color:var(--sub)}
.card .v{font-size:22px;color:var(--txt);margin-top:4px;font-weight:600}
.btns{display:flex;gap:6px} .btns button{flex:1;padding:8px;border:1px solid var(--brd);background:var(--bg);color:var(--txt);border-radius:6px;cursor:pointer;font-size:14px}
.btns button.on{background:var(--azul);color:#fff;border-color:var(--azul)}
table{border-collapse:collapse;font-size:12px;width:100%;margin-top:6px} th,td{border:1px solid var(--brd);padding:4px 6px;text-align:center} th{background:var(--card)}
tfoot td{font-weight:700;background:var(--card)}
h3{font-size:14px;margin:16px 0 0}
</style><script>__PLOTLYJS__</script></head><body>
__NAV__
<header><h1>__TITULO__</h1><div class="sub">__SUBTITULO__</div></header>
<div id="main"><div id="izq"><div id="grafico">__GRAFICO__</div>
<h3>Votos considerados en esta ronda</h3>__TABLA__
<h3>Trazabilidad por rondas (<span id="nom-est">media</span> de los votos vigentes en cada ronda)</h3>
<table id="tbl-hist"><thead><tr><th>Ronda</th><th>Evaluadores</th><th>Real a la fecha</th><th>Esperado resto del mes</th><th>Cierre de mes</th></tr></thead><tbody></tbody></table></div>
<aside id="der"><div class="card">Estadístico<div class="btns" style="margin-top:6px">
<button data-s="media" class="on">Media</button><button data-s="mediana">Mediana</button><button data-s="rango">Rango</button></div></div>
<div class="card">Real a la fecha (%)<div class="v" id="v-real"></div></div>
<div class="card">Esperado resto del mes (%)<div class="v" id="v-resto"></div></div>
<div class="card">Cierre de mes (%)<div class="v" id="v-cierre"></div></div></aside></div>
<script>
const D = __DATOS__;
const gd = document.getElementById('grafico').getElementsByClassName('plotly-graph-div')[0];
const $ = id => document.getElementById(id);
const r1 = v => (v===null||v===undefined) ? '' : String(Math.round(v*10)/10);
// valor de un estadístico {media,mediana,min,max} según el botón elegido (rango = "mín – máx")
const val = (e, s) => !e ? '—' : (s==='rango' ? (e.min===null ? '—' : r1(e.min)+' – '+r1(e.max)) : (e[s]===null ? '—' : r1(e[s])));
function mostrar(s){
  document.querySelectorAll('.btns button').forEach(b => b.classList.toggle('on', b.dataset.s===s));
  const todos = [].concat(...Object.values(D.grupos));
  Plotly.restyle(gd, {visible: false}, todos);
  Plotly.restyle(gd, {visible: true}, D.grupos[s]);
  $('nom-est').textContent = s;
  const real = D.real===null ? '—' : r1(D.real);
  $('v-real').textContent = real;
  $('v-resto').textContent = val(D.est.resto, s);
  $('v-cierre').textContent = val(D.est.cierre, s);
  // fila de totales de la tabla de votos
  const pie = $('pie-votos');
  D.semanas.forEach(k => pie.querySelector('[data-s="'+k+'"]').textContent = val(D.est.sem[k], s));
  pie.querySelector('[data-c="real"]').textContent = real;
  pie.querySelector('[data-c="resto"]').textContent = val(D.est.resto, s);
  pie.querySelector('[data-c="cierre"]').textContent = val(D.est.cierre, s);
  pie.querySelector('[data-c="nombre"]').textContent = {media:'Media',mediana:'Mediana',rango:'Rango (mín – máx)'}[s];
  // trazabilidad por rondas
  $('tbl-hist').querySelector('tbody').innerHTML = D.rondas.map(r =>
    '<tr><td>'+r.ronda+'</td><td>'+r.n+'</td><td>'+real+'</td><td>'+val(r.resto,s)+'</td><td>'+val(r.cierre,s)+'</td></tr>').join('');
}
document.querySelectorAll('.btns button').forEach(b => b.addEventListener('click', () => mostrar(b.dataset.s)));
window.addEventListener('load', () => setTimeout(() => Plotly.Plots.resize(gd), 50));
mostrar('media');
</script></body></html>
"""


def tabla_votos(vigentes, semanas, evaluadores, real_mean) -> str:
    cab = "".join(f"<th>S{s['n']}<br>{etiqueta_semana_corta(s)}</th>" for s in semanas)
    real = "—" if real_mean is None else f"{r1(real_mean)}"
    filas = []
    for v in vigentes:
        celdas = []
        for s in semanas:
            val = v["semanas"].get("S%d" % s["n"])
            celdas.append("<td>%s</td>" % ("" if val is None else r1(val)))
        filas.append(f"<tr><td>{html.escape(evaluadores.get(v['evaluador'], v['evaluador']))} "
                     f"({v['evaluador']})</td><td>{v['ronda']}</td>{''.join(celdas)}<td>{real}</td>"
                     f"<td>{r1(v['resto'])}</td><td>{r1(v['cierre'])}</td>"
                     f"<td>{v['fecha_hora']:%Y-%m-%d %H:%M}</td></tr>")
    pie = ("".join(f'<td data-s="S{s["n"]}"></td>' for s in semanas))
    return (f"<table><thead><tr><th>Evaluador</th><th>Ronda del voto</th>{cab}<th>Real a la fecha</th>"
            f"<th>Resto del mes</th><th>Cierre</th><th>Registrado</th></tr></thead><tbody>{''.join(filas)}</tbody>"
            f'<tfoot><tr id="pie-votos"><td data-c="nombre" colspan="2"></td>{pie}<td data-c="real"></td>'
            f'<td data-c="resto"></td><td data-c="cierre"></td><td></td></tr></tfoot></table>')


def escribir_consolidado(ruta, cuenca, ronda, vigentes, historial, serie, seg, orden_vars, fechas,
                         evaluadores, mes, real_mean, archivo):
    est = historial[ronda]
    fig, grupos = construir_consolidado(cuenca, vigentes, serie, seg, orden_vars, fechas, est, evaluadores)
    div = fig.to_html(full_html=False, include_plotlyjs=False,
                      config={"displaylogo": False, "responsive": True})
    datos = {
        "grupos": grupos, "real": real_mean, "semanas": [f"S{s['n']}" for s in fechas["semanas"]],
        "est": {"resto": est["resto"], "cierre": est["cierre"], "sem": est["sem"]},
        "rondas": [{"ronda": r, "n": e["n"], "resto": e["resto"], "cierre": e["cierre"]}
                   for r, e in sorted(historial.items()) if r <= ronda],
    }
    pagina = (PLANTILLA.replace("__TITULO_PAGINA__", html.escape(f"Consolidado {cuenca} R{ronda}"))
              .replace("__ESTILO_NAV__", ESTILO_NAV)
              .replace("__NAV__", "<!--NAV:%s:%d-->" % (archivo, ronda))
              .replace("__TITULO__", html.escape(f"Consolidado ronda {ronda} · {cuenca}"))
              .replace("__SUBTITULO__", html.escape(
                  f"Pronóstico {mes} · {len(vigentes)} evaluadores · "
                  f"horizonte {fechas['hor_ini']:%d/%m/%Y} – {fechas['hor_fin']:%d/%m/%Y}"))
              .replace("__TABLA__", tabla_votos(vigentes, fechas["semanas"], evaluadores, real_mean))
              .replace("__DATOS__", json.dumps(datos).replace("</", "<\\/"))
              .replace("__GRAFICO__", div)
              .replace("__PLOTLYJS__", plotly_js()))
    ruta.write_text(pagina, encoding="utf-8")


# --------------------------------------------------------------------------
# Principal
# --------------------------------------------------------------------------
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--aportes", default=BASE / ARCHIVO_APORTES, type=Path)
    ap.add_argument("--drivers", default=BASE / ARCHIVO_DRIVERS, type=Path)
    ap.add_argument("--pronostico", default=BASE / ARCHIVO_PRONOSTICO, type=Path)
    ap.add_argument("--consenso", default=BASE / ARCHIVO_CONSENSO, type=Path)
    ap.add_argument("--salida", default=SALIDAS, type=Path, help="carpeta donde está año/mes/fecha")
    ap.add_argument("--logs", default=BASE / "logs", type=Path)
    ap.add_argument("--fecha-corrida", default=None, help="AAAA-MM-DD (por defecto: fecha de realización del consenso)")
    ap.add_argument("--ronda", type=int, default=None, help="solo esta ronda (por defecto: todas las existentes)")
    args = ap.parse_args()

    run_id = uuid.uuid4().hex[:8]
    t0 = time.time()
    ruta_log = configurar_logging(args.logs, run_id)
    log.info("=" * 70)
    log.info("INICIO consolidar_votos run_id=%s", run_id)
    log.info("Usuario=%s host=%s SO=%s python=%s pandas=%s plotly=%s", getpass.getuser(),
             platform.node(), platform.platform(), platform.python_version(),
             pd.__version__, plotly.__version__)
    log.info("Argumentos: %s", {k: str(v) for k, v in vars(args).items()})
    try:
        for p in (args.aportes, args.drivers, args.pronostico, args.consenso):
            if not p.exists():
                raise FileNotFoundError(f"No se encontró el archivo Excel: {p} (debe estar en la carpeta del script)")
            st = p.stat()
            log.info("Entrada: %s | %d bytes | modificado %s | sha256=%s", p.resolve(), st.st_size,
                     datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds"), sha256(p))

        fechas = leer_fechas(args.consenso)
        evaluadores = leer_evaluadores(args.consenso)
        mes = f"{fechas['hor_ini']:%Y-%m}"
        fecha_corrida = (pd.Timestamp(args.fecha_corrida).normalize() if args.fecha_corrida
                         else fechas["consenso"])
        carpeta = carpeta_corrida(args.salida, fecha_corrida)
        carpeta_votos = carpeta / "votos"
        if not carpeta_votos.exists():
            raise FileNotFoundError(f"No existe la carpeta de votos: {carpeta_votos} "
                                    "(¿se corrió crear_graficos_consenso.py para esta fecha?)")
        (carpeta / CARPETA_CONSOLIDADO).mkdir(exist_ok=True)
        log.info("Carpeta de corrida: %s | mes del pronóstico: %s", carpeta.resolve(), mes)

        aportes = leer_aportes(args.aportes)
        ini, _ = definir_ventana(args.aportes, aportes)
        if ini >= fechas["hor_ini"]:
            ini = fechas["hor_ini"] - pd.Timedelta(days=30)
        cuencas = {c for c in aportes.columns if c != "Fecha"}
        ctx = cargar_drivers(args.drivers, args.pronostico)

        votos, rechazados = leer_votos(carpeta_votos, evaluadores, cuencas, mes)
        wb, ws, cols = cargar_registros(args.consenso)
        nuevos = agregar_votos(ws, cols, votos)
        todos = votos_del_mes(ws, cols, mes)
        log.info("Votos del mes %s en Registros: %d", mes, len(todos))

        # ---- cálculo por cuenca y ronda
        ahora = datetime.now().replace(microsecond=0)
        pendientes, filas_excel = [], []
        for cuenca in sorted({v["cuenca"] for v in todos}):
            if cuenca not in cuencas:
                log.warning("Cuenca '%s' de Registros no existe en Seguimiento_Aportes; se omite", cuenca)
                continue
            vc = [v for v in todos if v["cuenca"] == cuenca]
            rondas = sorted({v["ronda"] for v in vc})
            historial = {}
            vig_por_ronda = {}
            for r in rondas:
                vig_por_ronda[r] = votos_vigentes(vc, r)
                historial[r] = estadisticos_ronda(vig_por_ronda[r], fechas["semanas"])
            obs = real_a_la_fecha(aportes, cuenca, fechas["hor_ini"])
            real_mean = (sum(obs) / len(obs)) if obs else None
            serie = aportes.loc[(aportes["Fecha"] >= ini) & (aportes["Fecha"] <= fechas["hor_fin"]),
                                ["Fecha", cuenca]].dropna()
            seg, orden_vars = drivers_combinados(cuenca, ctx, ini, fechas["hor_fin"])
            for r in rondas:
                filas_excel += filas_consolidado(cuenca, r, vig_por_ronda[r], historial[r], real_mean, mes, ahora)
            for r in ([args.ronda] if args.ronda else rondas):
                if r not in vig_por_ronda:
                    log.warning("Cuenca '%s': no hay votos de la ronda %d; se omite", cuenca, r)
                    continue
                pendientes.append((cuenca, r, vig_por_ronda[r], historial, serie, seg, orden_vars,
                                   real_mean))

        # ---- Excel: Registros (votos nuevos) + Consolidado (resultados); una sola escritura
        n_nuevas, n_act = escribir_hoja_consolidado(wb, filas_excel) if filas_excel else (0, 0)
        log.info("Hoja Consolidado: %d filas nuevas, %d actualizadas", n_nuevas, n_act)
        if nuevos or n_nuevas or n_act:
            guardar_excel(wb, args.consenso, args.logs)
        else:
            log.info("Consenso.xlsx sin cambios; no se reescribe")

        # ---- gráficos Consolidado
        generados = []
        for cuenca, r, vig, historial, serie, seg, orden_vars, real_mean in pendientes:
            archivo = nombre_archivo(cuenca)
            registrar_cuenca(carpeta, archivo, cuenca)
            ruta = carpeta / CARPETA_CONSOLIDADO / f"{PREFIJO_CONSOLIDADO}_{archivo}_Ronda{r}.html"
            escribir_consolidado(ruta, cuenca, r, vig, historial, serie, seg, orden_vars, fechas,
                                 evaluadores, mes, real_mean, archivo)
            generados.append((archivo, r, ruta))
            log.info("Consolidado %s ronda %d -> %s | evaluadores=%s | media resto=%.1f cierre=%.1f | real=%s",
                     cuenca, r, ruta.name, [f"{v['evaluador']}(R{v['ronda']})" for v in vig],
                     historial[r]["resto"]["media"], historial[r]["cierre"]["media"],
                     "—" if real_mean is None else f"{real_mean:.1f}")
        # barra de navegación una vez que existen todos los archivos
        for f in (carpeta / CARPETA_CONSOLIDADO).glob(f"{PREFIJO_CONSOLIDADO}_*_Ronda*.html"):
            txt = f.read_text(encoding="utf-8")
            for archivo in cuencas_registradas(carpeta):
                for r in range(1, 51):
                    marca = "<!--NAV:%s:%d-->" % (archivo, r)
                    if marca in txt:
                        txt = txt.replace(marca, barra_navegacion(carpeta, archivo, "consolidado", r))
            f.write_text(txt, encoding="utf-8")
        escribir_indice(carpeta, fechas)

        log.info("RESUMEN: votos leídos=%d rechazados=%d nuevos en Registros=%d | Consolidado: +%d/~%d filas | gráficos=%d",
                 len(votos), len(rechazados), len(nuevos), n_nuevas, n_act, len(generados))
        log.info("FIN OK run_id=%s duración=%.2fs log=%s", run_id, time.time() - t0, ruta_log.resolve())
        print(f"\nVotos nuevos registrados: {len(nuevos)} | rechazados: {len(rechazados)} | "
              f"Consolidado Excel: +{n_nuevas}/~{n_act} filas | gráficos: {len(generados)}\n"
              f"Carpeta: {(carpeta / CARPETA_CONSOLIDADO).resolve()}")
        return 0
    except Exception:
        log.exception("FIN CON ERROR run_id=%s tras %.2fs", run_id, time.time() - t0)
        return 1


if __name__ == "__main__":
    sys.exit(main())
